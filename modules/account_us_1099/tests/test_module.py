# This file is part of Tryton.  The COPYRIGHT file at the top level of
# this repository contains the full copyright notices and license terms.
import datetime
import io
from decimal import Decimal

from trytond.model.exceptions import ValidationError
from trytond.modules.account.tests import create_chart, get_fiscalyear
from trytond.modules.company.tests import (
    CompanyTestMixin, create_company, set_company)
from trytond.pool import Pool
from trytond.tests.test_tryton import ModuleTestCase, with_transaction

from ..us_1099 import (
    FIRE_LAYOUTS, export_workbook, format_tin, truncate_tin)


def read_workbook(data):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data))
    sheet, = workbook.worksheets
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    return sheet.title, rows


class AccountUS1099TestCase(CompanyTestMixin, ModuleTestCase):
    'Test Account US 1099 module'
    module = 'account_us_1099'

    def setup(self, company):
        pool = Pool()
        Account = pool.get('account.account')
        FiscalYear = pool.get('account.fiscalyear')
        Journal = pool.get('account.journal')
        Party = pool.get('party.party')
        Identifier = pool.get('party.identifier')
        Country = pool.get('country.country')
        Subdivision = pool.get('country.subdivision')

        create_chart(company)
        for year in [2025, 2026]:
            fiscalyear = get_fiscalyear(
                company, today=datetime.date(year, 6, 1))
            fiscalyear.save()
            FiscalYear.create_period([fiscalyear])

        us = Country(name="United States", code='US')
        us.save()
        texas = Subdivision(
            name="Texas", code='US-TX', type='state', country=us)
        texas.save()

        Identifier(
            party=company.party, type='us_ein', code='123456789').save()
        address, = company.party.addresses
        address.street = '100 Main St'
        address.city = 'Rosenberg'
        address.postal_code = '77471'
        address.subdivision = texas
        address.country = us
        address.save()

        cash, = Account.search([('name', '=', "Main Cash")])
        cash.us_1099_payment = True
        cash.save()
        tax, = Account.search([('name', '=', "Main Tax")])
        card = Account(
            name="Credit Card", type=tax.type, company=company,
            parent=tax.parent)
        card.save()

        def party(name, box, tin_type=None, tin=None):
            party = Party(name=name, us_1099_box=box)
            if tin:
                party.identifiers = [{'type': tin_type, 'code': tin}]
            party.addresses = [{
                    'street': '5 Elm St\nSuite 2',
                    'city': 'Houston',
                    'postal_code': '77002',
                    'subdivision': texas.id,
                    'country': us.id,
                    }]
            party.save()
            return party

        self.contractor = party(
            "Jane Contractor", 'nec_1', 'us_ssn', '123-45-6789')
        self.landlord = party("Landlord LLC", 'misc_1', 'us_ein', '98-7654321')
        self.supplier = party("Supply Corp", None)
        self.cash = cash
        self.card = card
        self.payable, = Account.search([('type.payable', '=', True)])
        self.expense, = Account.search([('type.expense', '=', True)])
        self.journal, = Journal.search([('type', '=', 'cash')], limit=1)

    def move(self, company, date, lines):
        pool = Pool()
        Move = pool.get('account.move')
        Period = pool.get('account.period')
        move = Move(
            company=company, journal=self.journal, date=date,
            period=Period.find(company, date=date))
        move.lines = [{
                'account': a.id, 'party': p.id if p else None,
                'debit': Decimal(d), 'credit': Decimal(c),
                } for a, p, d, c in lines]
        move.save()
        Move.post([move])
        return move

    def pay(self, company, date, party, amount, account=None):
        "Pay the party from the account"
        return self.move(company, date, [
                (self.payable, party, amount, 0),
                (account or self.cash, None, 0, amount),
                ])

    @with_transaction()
    def test_compute(self):
        "Test compute 1099 amounts"
        pool = Pool()
        Amount = pool.get('account.us.1099')

        company = create_company()
        with set_company(company):
            self.setup(company)
            contractor, landlord = self.contractor, self.landlord
            # Invoice is not a payment
            self.move(company, datetime.date(2026, 2, 1), [
                    (self.expense, None, '3000', 0),
                    (self.payable, contractor, 0, '3000'),
                    ])
            self.pay(company, datetime.date(2026, 3, 1), contractor, '2500')
            # Paid by credit card
            self.pay(company, datetime.date(2026, 3, 2), contractor, '500',
                account=self.card)
            # Refund from the contractor
            self.move(company, datetime.date(2026, 4, 1), [
                    (self.cash, None, '200', 0),
                    (self.payable, contractor, 0, '200'),
                    ])
            # Previous year
            self.pay(company, datetime.date(2025, 12, 31), contractor, '700')
            self.pay(company, datetime.date(2026, 5, 1), landlord, '1500')
            self.pay(company, datetime.date(2026, 5, 1), self.supplier,
                '9000')

            amounts = Amount.compute(company, 2026)

            self.assertEqual(
                sorted((a.party.name, a.box, a.computed_amount)
                    for a in amounts), [
                    ("Jane Contractor", 'nec_1', Decimal('2300.00')),
                    ("Landlord LLC", 'misc_1', Decimal('1500.00')),
                    ])
            contractor_amount, = [
                a for a in amounts if a.party == contractor]
            landlord_amount, = [a for a in amounts if a.party == landlord]
            self.assertTrue(contractor_amount.reportable)
            self.assertFalse(landlord_amount.reportable)
            self.assertEqual(contractor_amount.form, 'nec')

            amounts_2025 = Amount.compute(company, 2025)
            self.assertEqual(
                [(a.party.name, a.computed_amount, a.reportable)
                    for a in amounts_2025],
                [("Jane Contractor", Decimal('700.00'), True)])

            # Adjustment is kept when computing again
            landlord_amount.adjustment = Decimal('600')
            landlord_amount.save()
            self.assertTrue(landlord_amount.reportable)
            self.pay(company, datetime.date(2026, 6, 1), landlord, '100')
            amounts = Amount.compute(company, 2026)
            landlord_amount, = [a for a in amounts if a.party == landlord]
            self.assertEqual(
                landlord_amount.computed_amount, Decimal('1600.00'))
            self.assertEqual(landlord_amount.amount, Decimal('2200.00'))
            self.assertEqual(len(amounts), 2)

            forms = Amount.forms(amounts)
            self.assertEqual(
                sorted((f.form, f.party.name, f.boxes) for f in forms), [
                    ('misc', "Landlord LLC", {'1': Decimal('2200.00')}),
                    ('nec', "Jane Contractor", {'1': Decimal('2300.00')}),
                    ])

    @with_transaction()
    def test_export(self):
        "Test export and recipient copies"
        pool = Pool()
        Amount = pool.get('account.us.1099')
        Export = pool.get('account.us.1099.export', type='wizard')
        RecipientCopy = pool.get('account.us.1099.recipient', type='report')

        company = create_company()
        with set_company(company):
            self.setup(company)
            self.pay(company, datetime.date(2026, 3, 1), self.contractor,
                '2500')
            self.pay(company, datetime.date(2026, 3, 1), self.landlord,
                '24000')
            amounts = Amount.compute(company, 2026)

            session_id, _, _ = Export.create()
            export = Export(session_id)
            export.start.company = company
            export.start.year = 2026
            self.assertEqual(export.transition_export(), 'result')
            self.assertEqual(export.result.nec_count, 1)
            self.assertEqual(export.result.misc_count, 1)
            self.assertEqual(
                export.result.nec_filename, 'F1099NEC-2026.xlsx')
            title, rows = read_workbook(export.result.nec_file)
            self.assertEqual(title, '1099-NEC Data')
            self.assertEqual(rows[0], [h for h, _ in FIRE_LAYOUTS['nec'][1]])
            row = dict(zip(rows[0], rows[1]))
            self.assertEqual(row['DO NOT USE'], None)
            self.assertEqual(row['Corrected (0, 1 or 2)'], 0)
            self.assertEqual(row['Payer TIN'], '12-3456789')
            self.assertEqual(row['Payer Address'], '100 Main St')
            self.assertEqual(row['Recipient TIN'], '123-45-6789')
            self.assertEqual(row['Recipient Address'], '5 Elm St|Suite 2')
            self.assertEqual(row['Recipient State'], 'TX')
            self.assertEqual(row['Recipient Zip'], '77002')
            self.assertEqual(
                row['Box 1 - Non Employee Compensation'], 2500)
            self.assertEqual(row['Box 4 - Federal Tax Withholding'], None)

            title, rows = read_workbook(export.result.misc_file)
            self.assertEqual(title, '1099-MISC Data')
            row = dict(zip(rows[0], rows[1]))
            self.assertEqual(row['Box 1\nRents'], 24000)
            self.assertEqual(row['Box 2\nRoyalties'], None)

            oext, content, _, _ = RecipientCopy.execute(
                [a.id for a in amounts], {})
            self.assertEqual(oext, 'pdf')
            self.assertTrue(content.startswith(b'%PDF'))

    @with_transaction()
    def test_export_missing_tin(self):
        "Test export without TIN"
        pool = Pool()
        Amount = pool.get('account.us.1099')
        Identifier = pool.get('party.identifier')

        company = create_company()
        with set_company(company):
            self.setup(company)
            self.pay(company, datetime.date(2026, 3, 1), self.contractor,
                '2500')
            Identifier.delete(list(self.contractor.identifiers))
            amounts = Amount.compute(company, 2026)
            form, = Amount.forms(amounts)
            with self.assertRaises(ValidationError):
                form.check()

    @with_transaction()
    def test_threshold(self):
        "Test thresholds by year"
        pool = Pool()
        Threshold = pool.get('account.us.1099.threshold')

        amounts_2025 = Threshold.get_amounts(2025)
        amounts_2026 = Threshold.get_amounts(2026)
        self.assertEqual(amounts_2025['nec_1'], Decimal('600'))
        self.assertEqual(amounts_2026['nec_1'], Decimal('2000'))
        self.assertEqual(amounts_2026['misc_2'], Decimal('10'))

    def test_tin(self):
        "Test TIN formatting"
        self.assertEqual(format_tin('us_ein', '123456789'), '12-3456789')
        self.assertEqual(format_tin('us_ssn', '123456789'), '123-45-6789')
        self.assertEqual(truncate_tin('123-45-6789'), '***-**-6789')
        self.assertEqual(truncate_tin('123'), '')

    def test_export_workbook_zip(self):
        "Test ZIP codes are kept as text"
        title, rows = read_workbook(export_workbook('nec', [{
                        'recipient_zip': '07001', 'box_1': Decimal('10'),
                        }]))
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row['Recipient Zip'], '07001')
        self.assertEqual(row['Box 1 - Non Employee Compensation'], 10)

del ModuleTestCase
